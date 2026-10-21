"""File parsing service: extract text from PDF and DOCX files, then chunk for embedding."""

import os
from typing import Optional
import pdfplumber
from docx import Document as DocxDocument


def parse_pdf(file_path: str) -> list[dict]:
    """Parse PDF and return list of {page, text} dicts."""
    pages = []
    with pdfplumber.open(file_path) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages.append({"page": i + 1, "text": text.strip()})
    return pages


def parse_docx(file_path: str) -> list[dict]:
    """Parse DOCX and return list of {page, text} dicts."""
    doc = DocxDocument(file_path)
    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())
    # Also extract table text
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                full_text.append(row_text)
    combined = "\n".join(full_text)
    return [{"page": 1, "text": combined}] if combined else []


def parse_file(file_path: str) -> list[dict]:
    """Parse a file and return list of {page, text} dicts based on file extension."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext in (".docx", ".doc"):
        return parse_docx(file_path)
    else:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()
        return [{"page": 1, "text": content}] if content.strip() else []


def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list[str]:
    """Split text into overlapping chunks of approximately chunk_size characters."""
    if len(text) <= chunk_size:
        return [text]
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        # Try to break at a sentence boundary
        if end < len(text):
            for sep in ["。", "！", "？", "\n", "；", " "]:
                last_sep = text.rfind(sep, start, end)
                if last_sep > start + chunk_size // 2:
                    end = last_sep + len(sep)
                    break
        chunks.append(text[start:end].strip())
        start = end - overlap
    return [c for c in chunks if c]
